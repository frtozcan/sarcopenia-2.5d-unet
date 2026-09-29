# -*- coding: utf-8 -*-
"""
Created on Tue Apr 14 16:25:46 2026

@author: FRTOZCAN
"""

"""
Sarkopeni - Otomatik Toplu DICOM Isleme (v4 - 2.5D Egitim Verisi)
==================================================================
YENİLİKLER:
  - SMI = tum kas dilimlerinin ORTALAMA alani / boy²
  - 2.5D egitim verisi: her kas dilimi icin
      [onceki_ct, merkez_ct, sonraki_ct] → 3 kanalli girdi
      merkez_maske                       → hedef cikti
  - Egitim/test bolmesi otomatik (%80 / %20)
  - Tum veriler numpy olarak kaydedilir

Kurulum:
  pip install pydicom numpy scikit-image pandas matplotlib tqdm openpyxl

Calistirma:
  python auto_batch2.py
"""

import os
import json
import numpy as np
import pandas as pd
import pydicom
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pathlib import Path
from skimage.draw import polygon
from tqdm import tqdm
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────
# AYARLAR
# ─────────────────────────────────────────
# Proje kok klasoru: SARKOPENI_DIR ortam degiskeniyle degistirilebilir
BASE_DIR = Path(os.environ.get("SARKOPENI_DIR", r"E:\sarkopeni"))

DATA_DIR   = BASE_DIR / "DATA"
OUTPUT_DIR = BASE_DIR / "output"

# 2.5D eğitim verisi çıktı klasörü
EGITIM_DIR  = OUTPUT_DIR / "egitim_25d"
TEST_BOLME  = 0.20   # %20 test, %80 eğitim
RANDOM_SEED = 42

# Klinik bilgiler (boy, kilo, yas, cinsiyet) bu Excel'den okunur.
# Excel'de bos olan alanlar icin DICOM basligindaki deger kullanilir.
KLINIK_EXCEL = BASE_DIR / "Hasta_bilgileri.xlsx"

# Egitim/test bolmesi bu dosyadan okunur (sutunlar: hasta_key, bolme).
# Dosya yoksa eski yontem kullanilir: seed=42 ile rastgele %80 / %20.
BOLME_CSV = EGITIM_DIR / "hasta_bolme.csv"

SMI_ESIK = {"e": 49.0, "k": 31.0, "m": 49.0, "f": 31.0}


# ═══════════════════════════════════════════
# DICOM'DAN KLİNİK VERİ
# ═══════════════════════════════════════════
def klinik_veri_cek(ds) -> dict:
    boy_m    = getattr(ds, "PatientSize",   None)
    boy_cm   = round(float(boy_m) * 100, 1) if boy_m else None
    kilo_raw = getattr(ds, "PatientWeight", None)
    kilo     = round(float(kilo_raw), 1) if kilo_raw else None
    yas_raw  = getattr(ds, "PatientAge", None)
    yas = None
    if yas_raw:
        try:
            yas = int("".join(filter(str.isdigit, str(yas_raw))))
        except Exception:
            pass
    cinsiyet_map = {"M": "E", "F": "K", "m": "E", "f": "K"}
    cinsiyet = cinsiyet_map.get(str(getattr(ds, "PatientSex", "")).strip(), None)
    bmi = round(kilo / (boy_cm / 100) ** 2, 1) if (boy_cm and kilo and boy_cm > 0) else None
    return {"boy_cm": boy_cm, "kilo_kg": kilo, "yas": yas, "cinsiyet": cinsiyet, "bmi": bmi}


# ═══════════════════════════════════════════
# SMI — ORTALAMA ALAN YÖNTEMİ
# ═══════════════════════════════════════════
def smi_hesapla_ortalama(kas_maske, spacing, boy_cm, cinsiyet):
    px_cm2 = (spacing[0] / 10.0) * (spacing[1] / 10.0)
    slice_alanlari_px = kas_maske.sum(axis=(1, 2))
    dolu_mask = slice_alanlari_px > 0
    dolu_px   = slice_alanlari_px[dolu_mask]

    if dolu_px.size == 0:
        return {"l3_idx": None, "kas_dilim_say": 0,
                "min_alan_cm2": None, "max_alan_cm2": None,
                "ort_alan_cm2": None, "smi": None,
                "smi_esik": None, "sarkopenik": None}

    l3_idx   = int(np.argmax(slice_alanlari_px))
    min_alan = round(float(dolu_px.min())  * px_cm2, 2)
    max_alan = round(float(dolu_px.max())  * px_cm2, 2)
    ort_alan = round(float(dolu_px.mean()) * px_cm2, 2)

    smi = smi_esik = sarkopenik = None
    if boy_cm and boy_cm > 0:
        smi      = round(ort_alan / (boy_cm / 100.0) ** 2, 2)
        smi_esik = SMI_ESIK.get((cinsiyet or "e").lower(), 49.0)
        sarkopenik = "EVET" if smi < smi_esik else "HAYIR"

    return {"l3_idx": l3_idx, "kas_dilim_say": int(dolu_mask.sum()),
            "min_alan_cm2": min_alan, "max_alan_cm2": max_alan,
            "ort_alan_cm2": ort_alan, "smi": smi,
            "smi_esik": smi_esik, "sarkopenik": sarkopenik}


# ═══════════════════════════════════════════
# 2.5D EĞİTİM VERİSİ HAZIRLAMA
# ═══════════════════════════════════════════
def egitim_verisi_kaydet(hasta_key: str,
                          ct_array: np.ndarray,
                          kas_maske: np.ndarray,
                          egitim_dir: Path,
                          test_dir: Path,
                          test_mi: bool):
    """
    Kas maskesi olan her dilim için 2.5D örnek oluşturur.

    Girdi (X): [onceki_ct, merkez_ct, sonraki_ct] → shape (H, W, 3)
    Hedef (Y): merkez dilimin kas maskesi          → shape (H, W, 1)

    HU penceresi: kas için -29 ~ 150, normalize [0, 1]

    Kaydedilen dosyalar:
        egitim/X/ veya test/X/  → {hasta_key}_z{idx:04d}_X.npy
        egitim/Y/ veya test/Y/  → {hasta_key}_z{idx:04d}_Y.npy
    """
    def hu_normalize(dilim):
        ct = np.clip(dilim.astype(np.float32), -29, 150)
        return (ct - (-29)) / (150 - (-29))

    hedef_dir = test_dir if test_mi else egitim_dir
    Z = ct_array.shape[0]
    slice_alanlari = kas_maske.sum(axis=(1, 2))
    kas_dilimleri  = np.where(slice_alanlari > 0)[0]

    kaydedilen = 0
    for idx in kas_dilimleri:
        onceki  = max(0, idx - 1)
        sonraki = min(Z - 1, idx + 1)

        # 3 kanallı CT girdisi (H, W, 3)
        kanal0 = hu_normalize(ct_array[onceki])[..., np.newaxis]
        kanal1 = hu_normalize(ct_array[idx]   )[..., np.newaxis]
        kanal2 = hu_normalize(ct_array[sonraki])[..., np.newaxis]
        X = np.concatenate([kanal0, kanal1, kanal2], axis=-1)  # (H, W, 3)

        # Hedef maske (H, W, 1)
        Y = kas_maske[idx][..., np.newaxis].astype(np.float32)  # (H, W, 1)

        dosya_adi = "{}_z{:04d}".format(hasta_key, idx)
        np.save(hedef_dir / "X" / "{}_X.npy".format(dosya_adi), X)
        np.save(hedef_dir / "Y" / "{}_Y.npy".format(dosya_adi), Y)
        kaydedilen += 1

    return kaydedilen


# ═══════════════════════════════════════════
# EXCEL
# ═══════════════════════════════════════════
def excel_kaydet(df: pd.DataFrame, dosya_yolu: Path, sayfa_adi="Sonuclar"):
    wb = Workbook()
    ws = wb.active
    ws.title = sayfa_adi

    BASLIK_DOLGU    = PatternFill("solid", fgColor="1F4E79")
    SARKOPENIK_DOLGU= PatternFill("solid", fgColor="FCE4D6")
    NORMAL_DOLGU    = PatternFill("solid", fgColor="E2EFDA")
    HATALI_DOLGU    = PatternFill("solid", fgColor="FCE4D6")
    EKSIK_DOLGU     = PatternFill("solid", fgColor="FFF2CC")
    BASLIK_YAZI     = Font(bold=True, color="FFFFFF", size=10)
    NORMAL_YAZI     = Font(size=10)
    ORTA = Alignment(horizontal="center", vertical="center")
    SOL  = Alignment(horizontal="left",   vertical="center")
    kenar = Side(style="thin", color="BFBFBF")
    KENAR = Border(left=kenar, right=kenar, top=kenar, bottom=kenar)

    SUTUNLAR = [
        ("No",                   "no",               5, "orta"),
        ("Hasta Adi",            "hasta_adi",        24, "sol"),
        ("Egitim/Test",          "egitim_test",      12, "orta"),
        ("Cinsiyet",             "cinsiyet",          9, "orta"),
        ("Yas",                  "yas",               7, "orta"),
        ("Boy (cm)",             "boy_cm",           10, "orta"),
        ("Kilo (kg)",            "kilo_kg",          10, "orta"),
        ("BMI",                  "bmi",               8, "orta"),
        ("CT Dilim Sayisi",      "ct_slice_sayisi",  15, "orta"),
        ("Pixel Spacing (mm)",   "pixel_spacing",    17, "orta"),
        ("Kesit Kalinligi (mm)", "kesit_kalinligi",  18, "orta"),
        ("Segmentler",           "segmentler",       30, "sol"),
        ("L3 Dilimi (merkez)",   "l3_dilim",         16, "orta"),
        ("Kas Dilim Sayisi",     "kas_dilim_say",    16, "orta"),
        ("2.5D Ornek Sayisi",    "ornek_sayisi",     16, "orta"),
        ("Min Kas Alani (cm2)",  "min_alan_cm2",     18, "orta"),
        ("Max Kas Alani (cm2)",  "max_alan_cm2",     18, "orta"),
        ("Ort Kas Alani (cm2)",  "ort_alan_cm2",     18, "orta"),
        ("Kas Hacim (cm3)",      "kas_hacim_cm3",    15, "orta"),
        ("Body Hacim (cm3)",     "body_hacim_cm3",   16, "orta"),
        ("SMI (cm2/m2)",         "smi",              13, "orta"),
        ("SMI Esigi",            "smi_esik",         11, "orta"),
        ("Sarkopenik",           "sarkopenik",       11, "orta"),
        ("Hata",                 "hata",             42, "sol"),
    ]

    ws.row_dimensions[1].height = 22
    for col_idx, (baslik, _, gen, _) in enumerate(SUTUNLAR, start=1):
        h = ws.cell(row=1, column=col_idx, value=baslik)
        h.fill = BASLIK_DOLGU; h.font = BASLIK_YAZI
        h.alignment = ORTA;    h.border = KENAR
        ws.column_dimensions[get_column_letter(col_idx)].width = gen

    for satir_no, (_, veri) in enumerate(df.iterrows(), start=2):
        ws.row_dimensions[satir_no].height = 17
        hata_var = bool(veri.get("hata") and
                        str(veri["hata"]).strip() not in ("", "nan", "None"))
        sark     = veri.get("sarkopenik")
        smi_var  = veri.get("smi") is not None and str(veri.get("smi")) != "nan"

        if hata_var:              dolgu = HATALI_DOLGU
        elif sark == "EVET":      dolgu = SARKOPENIK_DOLGU
        elif sark == "HAYIR":     dolgu = NORMAL_DOLGU
        elif not smi_var:         dolgu = EKSIK_DOLGU
        else:                     dolgu = NORMAL_DOLGU

        degerler = [satir_no - 1]
        for _, alan, _, _ in SUTUNLAR[1:]:
            degerler.append(veri.get(alan, ""))

        for col_idx, (deger, (_, _, _, hiz)) in enumerate(
                zip(degerler, SUTUNLAR), start=1):
            c = ws.cell(row=satir_no, column=col_idx, value=deger)
            c.font = NORMAL_YAZI; c.fill = dolgu; c.border = KENAR
            c.alignment = SOL if hiz == "sol" else ORTA

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(str(dosya_yolu))
    print("  Excel kaydedildi: {}".format(dosya_yolu))


# ═══════════════════════════════════════════
# YARDIMCI
# ═══════════════════════════════════════════
def anahtar(adi: str) -> str:
    return adi.strip().lower().replace(" ", "_")


def _bos_mu(deger) -> bool:
    return deger is None or (isinstance(deger, float) and np.isnan(deger)) \
        or str(deger).strip() in ("", "nan", "None")


def excel_klinik_yukle(yol: Path) -> dict:
    """
    Klinik Excel'ini okur -> {hasta_key: {boy_cm, kilo_kg, yas, cinsiyet}}.
    Beklenen sutunlar: 'Hasta Adi', 'Boy (cm)', 'Kilo (kg)', 'Yas', 'Cinsiyet'.
    """
    if not yol.exists():
        print("UYARI: Klinik Excel bulunamadi, DICOM degerleri kullanilacak:", yol)
        return {}
    df = pd.read_excel(yol)
    sonuc = {}
    for _, r in df.iterrows():
        if _bos_mu(r.get("Hasta Adi")):
            continue
        kayit = {}
        if not _bos_mu(r.get("Boy (cm)")):
            kayit["boy_cm"] = round(float(r["Boy (cm)"]), 1)
        if not _bos_mu(r.get("Kilo (kg)")):
            kayit["kilo_kg"] = round(float(r["Kilo (kg)"]), 1)
        if not _bos_mu(r.get("Yas")):
            kayit["yas"] = int(float(r["Yas"]))
        if not _bos_mu(r.get("Cinsiyet")):
            c = str(r["Cinsiyet"]).strip().upper()
            kayit["cinsiyet"] = {"M": "E", "F": "K"}.get(c, c)
        sonuc[anahtar(str(r["Hasta Adi"]))] = kayit
    print("Klinik Excel okundu: {} hasta ({})".format(len(sonuc), yol.name))
    return sonuc


def klinik_birlestir(dicom_klinik: dict, excel_kayit: dict) -> dict:
    """Excel degeri varsa onu, yoksa DICOM degerini kullanir; BMI'yi yeniden hesaplar."""
    k = dict(dicom_klinik)
    for alan in ("boy_cm", "kilo_kg", "yas", "cinsiyet"):
        if alan in excel_kayit:
            k[alan] = excel_kayit[alan]
    k["bmi"] = (round(k["kilo_kg"] / (k["boy_cm"] / 100) ** 2, 1)
                if (k.get("boy_cm") and k.get("kilo_kg")) else None)
    return k


def bolme_yukle(yol: Path) -> dict:
    """hasta_bolme.csv -> {hasta_key: 'EGITIM' | 'TEST'}; dosya yoksa bos dict."""
    if not yol.exists():
        return {}
    df = pd.read_csv(yol, encoding="utf-8-sig")
    return {anahtar(str(k)): str(b).strip().upper()
            for k, b in zip(df["hasta_key"], df["bolme"])}


# ═══════════════════════════════════════════
# CT YUKLEME
# ═══════════════════════════════════════════
def ct_yukle(klasor: Path):
    ct_dosyalar = sorted(
        [f for f in klasor.iterdir()
         if f.name.upper().startswith("CT") and f.suffix.lower() == ".dcm"],
        key=lambda f: float(
            pydicom.dcmread(str(f), stop_before_pixels=True).ImagePositionPatient[2])
    )
    if not ct_dosyalar:
        return None, None, None, None, None

    ref   = pydicom.dcmread(str(ct_dosyalar[0]))
    rows, cols, n = int(ref.Rows), int(ref.Columns), len(ct_dosyalar)
    pixel_spacing  = [float(x) for x in ref.PixelSpacing]
    image_position = [float(x) for x in ref.ImagePositionPatient]
    klinik = klinik_veri_cek(ref)

    ct_array = np.zeros((n, rows, cols), dtype=np.int16)
    slice_positions = []
    for i, f in enumerate(ct_dosyalar):
        ds = pydicom.dcmread(str(f))
        ct_array[i] = (ds.pixel_array
                       * float(getattr(ds, "RescaleSlope", 1))
                       + float(getattr(ds, "RescaleIntercept", 0)))
        slice_positions.append(float(ds.ImagePositionPatient[2]))

    slice_positions = np.array(slice_positions)
    z_spacing = float(np.abs(np.diff(slice_positions).mean())) if n > 1 else 5.0
    spacing   = (pixel_spacing[1], pixel_spacing[0], z_spacing)
    return ct_array, spacing, slice_positions, image_position, klinik


# ═══════════════════════════════════════════
# MASKE ÇIKARMA
# ═══════════════════════════════════════════
def maskeleri_cikar(struct_path, ct_array, spacing, slice_positions, image_position):
    ds = pydicom.dcmread(str(struct_path))
    maskeler = {}
    if not hasattr(ds, "ROIContourSequence"):
        return maskeler

    roi_isimleri = {roi.ROINumber: roi.ROIName.strip().lower()
                    for roi in ds.StructureSetROISequence}
    Z, H, W = ct_array.shape
    x0, y0  = image_position[0], image_position[1]
    dx, dy  = spacing[0], spacing[1]

    for roi_contour in ds.ROIContourSequence:
        roi_no   = roi_contour.ReferencedROINumber
        roi_isim = roi_isimleri.get(roi_no, "segment_{}".format(roi_no))
        maske    = np.zeros((Z, H, W), dtype=np.uint8)
        if not hasattr(roi_contour, "ContourSequence"):
            continue
        for contour in roi_contour.ContourSequence:
            if contour.ContourGeometricType != "CLOSED_PLANAR":
                continue
            pts   = np.array(contour.ContourData).reshape(-1, 3)
            z_idx = int(np.argmin(np.abs(slice_positions - pts[0, 2])))
            col_px = np.clip((pts[:, 0] - x0) / dx, 0, W - 1)
            row_px = np.clip((pts[:, 1] - y0) / dy, 0, H - 1)
            rr, cc = polygon(row_px, col_px, shape=(H, W))
            maske[z_idx, rr, cc] = 1
        maskeler[roi_isim] = maske
        print("    [{:20s}]  {:>8d} voxel".format(roi_isim, int(maske.sum())))
    return maskeler


# ═══════════════════════════════════════════
# L3 GÖRSELİ
# ═══════════════════════════════════════════
def gorsel_kaydet(img_dir, hasta_adi, ct_array, kas_maske, l3_idx):
    fig, axes = plt.subplots(1, 3, figsize=(15, 5))
    fig.suptitle(hasta_adi, fontsize=11)
    axes[0].imshow(ct_array[l3_idx],  cmap="gray", vmin=-200, vmax=200)
    axes[0].set_title("CT L3 merkez (z={})".format(l3_idx)); axes[0].axis("off")
    axes[1].imshow(kas_maske[l3_idx], cmap="Greens")
    axes[1].set_title("Kas maskesi"); axes[1].axis("off")
    axes[2].imshow(ct_array[l3_idx],  cmap="gray", vmin=-200, vmax=200)
    axes[2].imshow(kas_maske[l3_idx], cmap="Greens", alpha=0.45)
    axes[2].set_title("Birlesik"); axes[2].axis("off")
    plt.tight_layout()
    plt.savefig(img_dir / "{}_l3.png".format(anahtar(hasta_adi)), dpi=110)
    plt.close()


# ═══════════════════════════════════════════
# ANA DÖNGÜ
# ═══════════════════════════════════════════
def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    npy_dir = OUTPUT_DIR / "npy";       npy_dir.mkdir(exist_ok=True)
    img_dir = OUTPUT_DIR / "gorseller"; img_dir.mkdir(exist_ok=True)

    # 2.5D eğitim klasörleri
    egitim_x = EGITIM_DIR / "egitim" / "X"; egitim_x.mkdir(parents=True, exist_ok=True)
    egitim_y = EGITIM_DIR / "egitim" / "Y"; egitim_y.mkdir(parents=True, exist_ok=True)
    test_x   = EGITIM_DIR / "test"   / "X"; test_x.mkdir(parents=True,   exist_ok=True)
    test_y   = EGITIM_DIR / "test"   / "Y"; test_y.mkdir(parents=True,   exist_ok=True)

    egitim_dir = EGITIM_DIR / "egitim"
    test_dir   = EGITIM_DIR / "test"

    klasorler = sorted([d for d in DATA_DIR.iterdir() if d.is_dir()])
    toplam    = len(klasorler)
    print("Toplam hasta: {}".format(toplam))

    # Eğitim/test bölmesi — hasta bazlı (shuffle ile)
    klinik_excel = excel_klinik_yukle(KLINIK_EXCEL)
    bolme_haritasi = bolme_yukle(BOLME_CSV)

    if bolme_haritasi:
        print("Bolme dosyadan okundu:", BOLME_CSV)
        test_idx_set = {i for i, d in enumerate(klasorler)
                        if bolme_haritasi.get(anahtar(d.name)) == "TEST"}
        eksik = [d.name for d in klasorler
                 if anahtar(d.name) not in bolme_haritasi]
        if eksik:
            print("UYARI: bolme dosyasinda olmayan {} hasta EGITIM'e "
                  "alindi: {}".format(len(eksik), ", ".join(eksik)))
    else:
        print("Bolme dosyasi yok, rastgele bolme (seed={})".format(RANDOM_SEED))
        np.random.seed(RANDOM_SEED)
        sirali = np.random.permutation(toplam)
        test_idx_set = set(sirali[:int(toplam * TEST_BOLME)])
    print("Egitim: {}  |  Test: {}".format(
        toplam - len(test_idx_set), len(test_idx_set)))

    sonuclar, hatalar = [], []
    toplam_egitim_ornek = 0
    toplam_test_ornek   = 0

    for sira, klasor in enumerate(tqdm(klasorler, desc="Isleniyor")):
        hasta_adi = klasor.name
        hasta_key = anahtar(hasta_adi)
        test_mi   = sira in test_idx_set
        bolme_str = "TEST" if test_mi else "EGITIM"
        print("\n>>> {} [{}]".format(hasta_adi, bolme_str))

        try:
            # Struct
            struct_dosyalar = [
                f for f in klasor.iterdir()
                if f.suffix.lower() == ".dcm"
                and not f.name.upper().startswith("CT")
            ]
            if not struct_dosyalar:
                raise FileNotFoundError("struct .dcm bulunamadi")
            struct_path = max(struct_dosyalar, key=lambda f: f.stat().st_size)

            # CT + klinik
            ct_array, spacing, slice_pos, img_pos, klinik = ct_yukle(klasor)
            if ct_array is None:
                raise ValueError("CT*.dcm bulunamadi")
            if hasta_key in klinik_excel:
                klinik = klinik_birlestir(klinik, klinik_excel[hasta_key])
            else:
                print("  UYARI: hasta klinik Excel'de yok, DICOM degerleri kullanildi")
            print("  CT: {}  spacing=({:.4f},{:.4f},{:.2f})".format(
                ct_array.shape, *spacing))
            print("  Boy={}cm  Kilo={}kg  Yas={}  Cinsiyet={}  BMI={}".format(
                klinik["boy_cm"], klinik["kilo_kg"],
                klinik["yas"], klinik["cinsiyet"], klinik["bmi"]))

            # Maskeler
            maskeler = maskeleri_cikar(struct_path, ct_array,
                                       spacing, slice_pos, img_pos)
            if not maskeler:
                raise ValueError("Segment cikartilamadi")

            # Ham npy kaydet
            hasta_dir = npy_dir / hasta_key
            hasta_dir.mkdir(exist_ok=True)
            np.save(hasta_dir / "ct.npy",      ct_array)
            np.save(hasta_dir / "spacing.npy", np.array(spacing))

            segment_stats = {}
            for seg_isim, maske in maskeler.items():
                np.save(hasta_dir / "mask_{}.npy".format(seg_isim), maske)
                voxel = int(maske.sum())
                segment_stats[seg_isim] = {
                    "voxel_count": voxel,
                    "volume_cm3": round(
                        voxel * spacing[0] * spacing[1] * spacing[2] / 1000, 2)
                }

            # SMI (ortalama)
            kas_key  = next((k for k in maskeler if "kas" in k.lower()), None)
            body_key = next((k for k in maskeler if "body" in k.lower()), None)
            smi_sonu = {"l3_idx": None, "kas_dilim_say": None,
                        "min_alan_cm2": None, "max_alan_cm2": None,
                        "ort_alan_cm2": None, "smi": None,
                        "smi_esik": None, "sarkopenik": None}
            ornek_sayisi = 0

            if kas_key:
                kas_maske = maskeler[kas_key]
                smi_sonu  = smi_hesapla_ortalama(
                    kas_maske, spacing,
                    klinik["boy_cm"], klinik["cinsiyet"])

                if smi_sonu["l3_idx"] is not None:
                    # L3 merkez dilimi kaydet
                    l3 = smi_sonu["l3_idx"]
                    np.save(hasta_dir / "l3_ct_slice.npy",   ct_array[l3])
                    np.save(hasta_dir / "l3_mask_slice.npy", kas_maske[l3])
                    gorsel_kaydet(img_dir, hasta_adi,
                                  ct_array, kas_maske, l3)

                    # ── 2.5D EĞİTİM VERİSİ ──────────────────────────
                    # Hasta bolme degistirdiyse diger klasordeki eski
                    # orneklerini sil (ayni hasta iki sette olmasin)
                    diger_dir = egitim_dir if test_mi else test_dir
                    for xy in ("X", "Y"):
                        for eski in (diger_dir / xy).glob(
                                "{}_z*_{}.npy".format(hasta_key, xy)):
                            eski.unlink()

                    ornek_sayisi = egitim_verisi_kaydet(
                        hasta_key, ct_array, kas_maske,
                        egitim_dir, test_dir, test_mi)

                    if test_mi:
                        toplam_test_ornek   += ornek_sayisi
                    else:
                        toplam_egitim_ornek += ornek_sayisi

                    print("  Kas dilimleri: {}  |  2.5D ornek: {}  |  "
                          "Ort.alan={} cm2  |  SMI={}  |  {}".format(
                        smi_sonu["kas_dilim_say"], ornek_sayisi,
                        smi_sonu["ort_alan_cm2"],
                        smi_sonu["smi"],
                        smi_sonu["sarkopenik"] or "boy eksik"))
            else:
                print("  UYARI: kas segmenti yok →", list(maskeler.keys()))

            kas_hacim  = segment_stats[kas_key]["volume_cm3"]  if kas_key  else None
            body_hacim = segment_stats[body_key]["volume_cm3"] if body_key else None

            meta = {
                "hasta_adi": hasta_adi, "hasta_key": hasta_key,
                "egitim_test": bolme_str, **klinik,
                "spacing_mm": list(spacing),
                "pixel_spacing": round(spacing[0], 4),
                "kesit_kalinligi": round(spacing[2], 2),
                "ct_shape": list(ct_array.shape),
                "segments": list(maskeler.keys()),
                "segment_stats": segment_stats,
                **smi_sonu,
                "ornek_sayisi_25d": ornek_sayisi,
                "kas_hacim_cm3": kas_hacim,
                "body_hacim_cm3": body_hacim,
            }
            with open(hasta_dir / "meta.json", "w", encoding="utf-8") as f:
                json.dump(meta, f, ensure_ascii=False, indent=2)

            sonuclar.append({
                "hasta_adi":       hasta_adi,
                "egitim_test":     bolme_str,
                "cinsiyet":        klinik["cinsiyet"],
                "yas":             klinik["yas"],
                "boy_cm":          klinik["boy_cm"],
                "kilo_kg":         klinik["kilo_kg"],
                "bmi":             klinik["bmi"],
                "ct_slice_sayisi": ct_array.shape[0],
                "pixel_spacing":   round(spacing[0], 4),
                "kesit_kalinligi": round(spacing[2], 2),
                "segmentler":      ", ".join(maskeler.keys()),
                "l3_dilim":        smi_sonu["l3_idx"],
                "kas_dilim_say":   smi_sonu["kas_dilim_say"],
                "ornek_sayisi":    ornek_sayisi,
                "min_alan_cm2":    smi_sonu["min_alan_cm2"],
                "max_alan_cm2":    smi_sonu["max_alan_cm2"],
                "ort_alan_cm2":    smi_sonu["ort_alan_cm2"],
                "kas_hacim_cm3":   kas_hacim,
                "body_hacim_cm3":  body_hacim,
                "smi":             smi_sonu["smi"],
                "smi_esik":        smi_sonu["smi_esik"],
                "sarkopenik":      smi_sonu["sarkopenik"],
                "hata":            None,
            })
            print("  OK")

        except Exception as e:
            print("  HATA:", e)
            hatalar.append({"hasta_adi": hasta_adi, "hata": str(e)})
            sonuclar.append({
                "hasta_adi": hasta_adi, "egitim_test": bolme_str,
                "cinsiyet": None, "yas": None, "boy_cm": None,
                "kilo_kg": None, "bmi": None, "ct_slice_sayisi": None,
                "pixel_spacing": None, "kesit_kalinligi": None,
                "segmentler": None, "l3_dilim": None, "kas_dilim_say": None,
                "ornek_sayisi": None, "min_alan_cm2": None,
                "max_alan_cm2": None, "ort_alan_cm2": None,
                "kas_hacim_cm3": None, "body_hacim_cm3": None,
                "smi": None, "smi_esik": None, "sarkopenik": None,
                "hata": str(e),
            })

    # Excel
    df = pd.DataFrame(sonuclar)
    excel_kaydet(df, OUTPUT_DIR / "isleme_sonuclari.xlsx")

    # Özet
    df_smi = df[df["smi"].notna()]
    print("\n" + "=" * 65)
    print("TAMAMLANDI")
    print("  Toplam hasta              : {}".format(len(sonuclar)))
    print("  Basarili                  : {}".format(len(sonuclar) - len(hatalar)))
    print("  Hatali                    : {}".format(len(hatalar)))
    print("  SMI hesaplanan            : {}".format(len(df_smi)))
    if len(df_smi) > 0:
        sark_n = (df_smi["sarkopenik"] == "EVET").sum()
        print("  Sarkopenik                : {} ({:.1f}%)".format(
            sark_n, 100 * sark_n / len(df_smi)))
        print("  Ortalama SMI              : {:.1f} cm2/m2".format(
            df_smi["smi"].mean()))
    print("  --- 2.5D Egitim Verisi ---")
    print("  Egitim ornegi (toplam)    : {}".format(toplam_egitim_ornek))
    print("  Test ornegi (toplam)      : {}".format(toplam_test_ornek))
    print("  Egitim klasoru            : {}".format(EGITIM_DIR / "egitim"))
    print("  Test klasoru              : {}".format(EGITIM_DIR / "test"))
    print("  Dosya yapisi              : X/{hasta}_z{idx}_X.npy (H,W,3)")
    print("                              Y/{hasta}_z{idx}_Y.npy (H,W,1)")
    print("=" * 65)

    if hatalar:
        print("\nHATALI HASTALAR:")
        for h in hatalar:
            print("  {:30s}  {}".format(h["hasta_adi"], h["hata"]))

    return df


if __name__ == "__main__":
    main()